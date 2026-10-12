from io import BytesIO
import unittest
from fastapi import UploadFile
from langchain_core.documents import Document
from chunking import split_documents
from ingestion import load_chunks


class ChunkingTests(unittest.TestCase):
    def test_headings_separate_short_sections(self):
        chunks = split_documents([Document(page_content='# 设备手册\n## 焊接\n检查气体。\n\n## 冷却\n检查冷却水。')])
        self.assertEqual(len(chunks), 2)
        self.assertEqual(chunks[0].metadata['section_path'], '设备手册 > 焊接')
        self.assertNotIn('冷却水', chunks[0].page_content)
        self.assertNotIn('检查气体', chunks[1].page_content)

    def test_long_section_repeats_ancestry_and_respects_size(self):
        text = '# 设备手册\n## 焊接\n' + '检查保护气体纯度，检查聚焦镜位置。' * 80
        chunks = split_documents([Document(page_content=text)], chunk_size=120, chunk_overlap=20)
        self.assertGreater(len(chunks), 1)
        for chunk in chunks:
            self.assertTrue(chunk.page_content.startswith('设备手册 > 焊接\n\n'))
            self.assertLessEqual(len(chunk.page_content), 120)
            body = chunk.page_content.split('\n\n', 1)[1]
            offset = chunk.metadata['start_index']
            self.assertEqual(text[offset:offset + len(body)], body)

    def test_code_fence_headings_are_not_sections(self):
        chunks = split_documents([Document(page_content='# 示例\n```python\n# comment\nprint(1)\n```\n正常正文')])
        self.assertEqual(len(chunks), 1)
        self.assertEqual(chunks[0].metadata['section_path'], '示例')
        self.assertIn('# comment', chunks[0].page_content)

    def test_chinese_and_numbered_hierarchy(self):
        chunks = split_documents([Document(page_content='第一章 维护\n1.1 焊接设备\n检查保护气体\n1.2 冷却设备\n检查冷却水')])
        self.assertEqual(chunks[0].metadata['section_path'], '第一章 维护 > 1.1 焊接设备')
        self.assertEqual(chunks[1].metadata['section_path'], '第一章 维护 > 1.2 冷却设备')

    def test_recursive_mode_and_paragraph_fallback(self):
        text = '# 手册\n\n第一段正文。\n\n第二段正文。'
        chunks = split_documents([Document(page_content=text)], strategy='recursive')
        self.assertIn('# 手册', chunks[0].page_content)
        self.assertEqual(chunks[0].metadata['section_path'], '')
        self.assertEqual(chunks[0].metadata['split_strategy'], 'recursive')
        plain = split_documents([Document(page_content='检查保护气体。\n\n检查冷却水。')])
        self.assertEqual(plain[0].page_content, '检查保护气体。\n\n检查冷却水。')

    def test_pdf_page_metadata_stays_with_chunks(self):
        chunks = split_documents([Document(page_content='# 焊接\n检查保护气体', metadata={'page': 2}),
                                  Document(page_content='# 冷却\n检查冷却水', metadata={'page': 3})])
        self.assertEqual([chunk.metadata['page'] for chunk in chunks], [2, 3])

    def test_invalid_strategy_is_422(self):
        from fastapi import HTTPException
        with self.assertRaises(HTTPException) as error:
            load_chunks(UploadFile(filename='a.txt', file=BytesIO(b'text')), strategy='unknown')
        self.assertEqual(error.exception.status_code, 422)

    def test_pdf_continuation_keeps_section_path(self):
        chunks = split_documents([Document(page_content='# 焊接\n检查保护气体', metadata={'filename': 'a.pdf', 'page': 1}),
                                  Document(page_content='继续检查聚焦镜', metadata={'filename': 'a.pdf', 'page': 2})])
        self.assertEqual(chunks[1].metadata['section_path'], '焊接')
        self.assertEqual(chunks[1].metadata['page'], 2)

    def test_docx_heading_styles_and_table_text(self):
        from docx import Document as WordDocument
        doc = WordDocument()
        doc.add_heading('设备手册', level=1)
        doc.add_heading('焊接', level=2)
        doc.add_paragraph('检查保护气体')
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = '型号'
        table.cell(0, 1).text = '设备A'
        doc.add_heading('冷却', level=2)
        doc.add_paragraph('检查冷却水')
        stream = BytesIO()
        doc.save(stream)
        stream.seek(0)
        chunks, _ = load_chunks(UploadFile(filename='a.docx', file=stream))
        self.assertEqual(len(chunks), 2)
        self.assertEqual(chunks[0].metadata['section_path'], '设备手册 > 焊接')
        self.assertIn('型号 | 设备A', chunks[0].page_content)
        self.assertNotIn('冷却水', chunks[0].page_content)


if __name__ == '__main__':
    unittest.main()
