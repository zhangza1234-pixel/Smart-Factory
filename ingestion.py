"""Document loading with bounded size, page metadata and stable chunk identity."""
from hashlib import sha256
import json
import re
from pathlib import Path
import tempfile
from fastapi import HTTPException
from langchain_core.documents import Document
from langchain_community.document_loaders import PyPDFLoader
from chunking import split_documents


def load_docx(path):
    """Translate Word heading styles into structural markers before splitting."""
    from docx import Document as WordDocument
    from docx.text.paragraph import Paragraph
    lines = []
    for block in WordDocument(path).iter_inner_content():
        if isinstance(block, Paragraph):
            style = block.style.name if block.style else ""
            match = re.fullmatch(r"(?:Heading|标题)\s*([1-6])", style, re.IGNORECASE)
            text = block.text
            if match and text.strip():
                text = "#" * int(match[1]) + " " + text
            lines.append(text)
        else:
            lines.extend(" | ".join(cell.text for cell in row.cells) for row in block.rows)
    return [Document(page_content="\n\n".join(lines))]


def load_chunks(file, chunk_size=800, chunk_overlap=120, max_bytes=20 * 1024 * 1024, strategy="structured"):
    filename = Path(file.filename or "").name
    extension = Path(filename).suffix.lower()
    if extension not in {".pdf", ".txt", ".md", ".docx"}:
        raise HTTPException(415, "支持 PDF、TXT、MD、DOCX 文件")
    if chunk_size < 50 or not 0 <= chunk_overlap < chunk_size or strategy not in {"structured", "recursive"}:
        raise HTTPException(422, "要求 size >= 50、0 <= overlap < size，策略为 structured 或 recursive")
    content = file.file.read(max_bytes + 1)
    if len(content) > max_bytes:
        raise HTTPException(413, "文档超过上传大小限制")
    if not content:
        raise HTTPException(422, "文件内容为空")
    try:
        if extension in {".txt", ".md"}:
            try:
                text = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = content.decode("gb18030")
            docs = [Document(page_content=text)]
        else:
            with tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / ("document" + extension)
                path.write_bytes(content)
                docs = PyPDFLoader(str(path)).load() if extension == ".pdf" else load_docx(path)
        for doc in docs:
            page = doc.metadata.get("page")
            doc.metadata = {"filename": filename, "page": int(page) + 1 if page is not None else 0}
        chunks = split_documents(docs, chunk_size, chunk_overlap, strategy)
        if not chunks:
            raise HTTPException(422, "未提取到文本；扫描 PDF 请先进行 OCR")
        return chunks, sha256(content).hexdigest()
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, "文档解析失败，请检查文件格式或编码") from exc


def prepare_chunks(chunks, content_hash, metadata):
    identity = json.dumps({"content": content_hash, "metadata": metadata,
                           "chunks": [(chunk.page_content, chunk.metadata) for chunk in chunks]},
                          ensure_ascii=False, sort_keys=True)
    document_id = sha256(identity.encode()).hexdigest()
    ids, texts, metadatas = [], [], []
    for index, chunk in enumerate(chunks):
        ids.append(f"{document_id}:{index}")
        texts.append(chunk.page_content)
        metadatas.append({**metadata, **chunk.metadata, "document_id": document_id,
                          "content_hash": content_hash, "chunk_index": index})
    return document_id, ids, texts, metadatas
